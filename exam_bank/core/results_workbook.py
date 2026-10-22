"""The results workbook: one row per exam ATTEMPT, as the Content Manager's
webhook writes it to the results Sheet and as Google downloads it to .xlsx.

The workbook has three kinds of tab:

- **per-course attempt tabs** ("DI Practitioner", "BA Practitioner", ...): one
  row per attempt - attempt id, candidate id, course id, submitted at, score,
  total, percentage, PASS/FAIL, role, company, partner id;
- **All results**: the same rows for every course together (so each attempt
  appears twice in the workbook; it is imported once);
- **Detail**: one row per question per attempt - what `item_analysis` reads.

A candidate who sits an exam again has a second row under the same candidate
id and a different attempt id. Both are kept: the attempt number is worked
out from the order of a candidate's submissions to that course, never stored,
so importing the older attempt later still numbers them correctly.

Privacy: names, emails and free-text notes are never read. A candidate is the
opaque `cand_...` id the exam assigned, which is enough to tell "the same
person again" without saying who.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Iterable, Optional

from . import item_analysis as ia

#: Attempt-tab headings (lower-cased) -> the field they fill.
ATTEMPT_FIELDS = {
    "attempt id": "attempt_id",
    "candidate id": "candidate_id",
    "course id": "course_id",
    "role": "role",
    "submitted at": "submitted_at",
    "company": "company",
    "score": "score",
    "total": "total",
    "percentage": "percentage",
    "passed": "passed",
    "partner id": "partner_id",
}
#: The tab that repeats every course's rows; the per-course tab wins on a clash.
ALL_TAB = "all results"


class WorkbookError(ia.ResultsImportError):
    """The file is not the results workbook; the message says what is wrong."""


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _number(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _stamp(v: Any) -> str:
    """ISO text for a date cell; text cells are kept as written."""
    if isinstance(v, datetime):
        return v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    return _text(v)


def day_of(stamp: str) -> str:
    """The YYYY-MM-DD of a stored timestamp ('' when it has none)."""
    s = (stamp or "").strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-" and s[:4].isdigit():
        return s[:10]
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _passed(v: Any) -> Optional[int]:
    t = _text(v).upper()
    if t in ("PASS", "PASSED", "Y", "YES", "TRUE", "1"):
        return 1
    if t in ("FAIL", "FAILED", "N", "NO", "FALSE", "0"):
        return 0
    return None


def parse_workbook(data: bytes) -> dict[str, Any]:
    """Attempts and answers from every tab of the workbook.

    Returns ``{"attempts": [...], "answers": [...], "sheets": [...]}``. Raises
    WorkbookError when the file is not an .xlsx or has neither kind of tab.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises a zoo of types for a non-workbook
        raise WorkbookError("This is not an Excel workbook (.xlsx).") from e

    attempts: dict[str, dict[str, Any]] = {}
    answers: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []
    try:
        # Per-course tabs first, so an attempt keeps the tab it is listed under.
        ordered = sorted(wb.worksheets, key=lambda ws: ws.title.strip().lower() == ALL_TAB)
        for ws in ordered:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            keys = [_text(h).lower() for h in (header or ())]
            info = {"name": ws.title, "kind": "ignored", "rows": 0}
            if "attempt id" in keys and "candidate id" in keys:
                info["kind"] = "attempts"
                idx = {ATTEMPT_FIELDS[k]: i for i, k in enumerate(keys) if k in ATTEMPT_FIELDS}
                tab_is_all = ws.title.strip().lower() == ALL_TAB
                for line in rows:
                    cell = lambda f: line[idx[f]] if f in idx and idx[f] < len(line) else None
                    attempt_id = _text(cell("attempt_id"))
                    if not attempt_id:
                        continue
                    info["rows"] += 1
                    if attempt_id in attempts:
                        continue
                    course_id = _text(cell("course_id"))
                    if not course_id and not tab_is_all:
                        course_id = ws.title.strip()
                    score, total, pct = _number(cell("score")), _number(cell("total")), _number(cell("percentage"))
                    if pct is None and score is not None and total:
                        pct = round(score / total, 4)
                    attempts[attempt_id] = {
                        "attempt_id": attempt_id,
                        "candidate_id": _text(cell("candidate_id")),
                        "course_id": course_id,
                        "course_tab": "" if tab_is_all else ws.title.strip(),
                        "role": _text(cell("role")),
                        "submitted_at": _stamp(cell("submitted_at")),
                        "company": _text(cell("company")),
                        "score": score,
                        "total": total,
                        "percentage": pct,
                        "passed": _passed(cell("passed")),
                        "partner_id": _text(cell("partner_id")),
                    }
            elif all(k in keys for k in ia.REQUIRED):
                info["kind"] = "answers"
                answers.extend(_detail_rows(keys, rows, info))
            sheets.append(info)
    finally:
        wb.close()

    if not attempts and not answers:
        raise WorkbookError(
            "This is not the results workbook: no tab has the attempt columns "
            '("Attempt ID", "Candidate ID", ...) or the Detail columns '
            '("Attempt ID", "Question ID", "Was correct?", ...).'
        )
    return {"attempts": list(attempts.values()), "answers": answers, "sheets": sheets}


def _detail_rows(keys: list[str], rows: Iterable[tuple], info: dict[str, Any]) -> list[dict[str, Any]]:
    """The Detail tab's rows, by the same rules as the CSV import."""
    fields = {**ia.REQUIRED, **{k: v for k, v in ia.OPTIONAL.items() if k in keys}}
    idx = {name: keys.index(h) for h, name in fields.items()}
    out = []
    for line in rows:
        row = {n: (_stamp(line[i]) if n == "submitted_at" else _text(line[i])) if i < len(line) else ""
               for n, i in idx.items()}
        if not row["attempt_id"] or not row["question_id"]:
            continue
        info["rows"] += 1
        mark = row["was_correct"].upper()
        if mark == "-":
            continue
        row["was_correct"] = 1 if mark in ("Y", "YES", "TRUE", "1") else 0
        row.setdefault("submitted_at", "")
        out.append(row)
    return out


def import_attempts(conn, attempts: Iterable[dict[str, Any]]) -> dict[str, int]:
    """Upsert by attempt id: importing the same workbook twice changes nothing."""
    ia.ensure_schema(conn)
    rows = list(attempts)
    before = conn.execute("SELECT COUNT(*) FROM attempts").fetchone()[0]
    conn.executemany(
        "INSERT INTO attempts (attempt_id, candidate_id, course_id, course_tab, role, submitted_at, "
        "company, score, total, percentage, passed, partner_id) VALUES "
        "(:attempt_id, :candidate_id, :course_id, :course_tab, :role, :submitted_at, :company, "
        ":score, :total, :percentage, :passed, :partner_id) "
        "ON CONFLICT(attempt_id) DO UPDATE SET candidate_id=excluded.candidate_id, "
        "course_id=excluded.course_id, course_tab=excluded.course_tab, role=excluded.role, "
        "submitted_at=excluded.submitted_at, company=excluded.company, score=excluded.score, "
        "total=excluded.total, percentage=excluded.percentage, passed=excluded.passed, "
        "partner_id=excluded.partner_id",
        rows,
    )
    conn.commit()
    after = conn.execute("SELECT COUNT(*) FROM attempts").fetchone()[0]
    return {"attempts": len(rows), "addedAttempts": after - before, "totalAttempts": after}


def import_workbook(conn, data: bytes) -> dict[str, Any]:
    """Import every tab: the attempts, and the Detail answers when it is there."""
    parsed = parse_workbook(data)
    result = ia.import_rows(conn, parsed["answers"]) if parsed["answers"] else {
        "rows": 0, "added": 0, "attempts": 0, "courses": [],
        "totalResponses": conn.execute("SELECT COUNT(*) FROM responses").fetchone()[0]
        if _has_table(conn, "responses") else 0,
    }
    mine = import_attempts(conn, parsed["attempts"])
    courses = sorted({a["course_id"] for a in parsed["attempts"] if a["course_id"]} | set(result["courses"]))
    return {**result, **mine, "answers": result["rows"], "attempts": mine["attempts"],
            "answerAttempts": result["attempts"], "courses": courses, "sheets": parsed["sheets"]}


def _has_table(conn, name: str) -> bool:
    return bool(conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,)).fetchone())


# -- reading back, filtered --------------------------------------------------

_COLUMNS = ("attempt_id", "candidate_id", "course_id", "course_tab", "role", "submitted_at",
            "company", "score", "total", "percentage", "passed", "partner_id")


def attempt_index(conn) -> list[dict[str, Any]]:
    """Every known attempt with its number, oldest first.

    An attempt that only exists as Detail answers (a CSV import) is listed too,
    with what the answers say (course, date) and no candidate or score.
    """
    ia.ensure_schema(conn)
    found = {}
    for r in conn.execute(f"SELECT {', '.join(_COLUMNS)} FROM attempts"):
        found[r[0]] = dict(zip(_COLUMNS, r))
    for aid, course, stamp in conn.execute(
        "SELECT attempt_id, course_id, MIN(submitted_at) FROM responses GROUP BY attempt_id, course_id"
    ):
        found.setdefault(aid, {
            "attempt_id": aid, "candidate_id": "", "course_id": course or "", "course_tab": "",
            "role": "", "submitted_at": stamp or "", "company": "", "score": None, "total": None,
            "percentage": None, "passed": None, "partner_id": "",
        })
    ordered = sorted(found.values(), key=lambda a: (a["submitted_at"] or "", a["attempt_id"]))
    seen: dict[tuple[str, str], int] = defaultdict(int)
    for a in ordered:
        a["day"] = day_of(a["submitted_at"])
        if a["candidate_id"]:
            seen[(a["candidate_id"], a["course_id"])] += 1
            a["attempt_no"] = seen[(a["candidate_id"], a["course_id"])]
        else:
            a["attempt_no"] = None
    return ordered


class Filters:
    """What the Results view narrows by. Empty means no restriction."""

    def __init__(self, course: str = "", candidate: str = "", passed: str = "", attempt_no: int = 0,
                 date_from: str = "", date_to: str = ""):
        self.course = course.strip()
        self.candidate = candidate.strip()
        self.passed = passed.strip().lower()
        self.attempt_no = attempt_no or 0
        self.date_from = date_from.strip()
        self.date_to = date_to.strip()
        if self.passed not in ("", "pass", "fail"):
            raise ValueError('passed must be "pass" or "fail"')
        for d in (self.date_from, self.date_to):
            if d:
                date.fromisoformat(d)

    @property
    def active(self) -> bool:
        return any((self.course, self.candidate, self.passed, self.attempt_no, self.date_from, self.date_to))

    def matches(self, a: dict[str, Any]) -> bool:
        if self.course and a["course_id"] != self.course:
            return False
        if self.candidate and self.candidate.lower() not in a["candidate_id"].lower():
            return False
        if self.passed and a["passed"] != (1 if self.passed == "pass" else 0):
            return False
        if self.attempt_no and a["attempt_no"] != self.attempt_no:
            return False
        if self.date_from or self.date_to:
            if not a["day"]:
                return False
            if self.date_from and a["day"] < self.date_from:
                return False
            if self.date_to and a["day"] > self.date_to:
                return False
        return True


def matching_attempt_ids(conn, flt: Filters) -> Optional[set[str]]:
    """The attempt ids the filters let through; None when nothing is filtered."""
    if not flt.active:
        return None
    return {a["attempt_id"] for a in attempt_index(conn) if flt.matches(a)}


def list_attempts(conn, flt: Filters, course_names: Optional[dict[str, str]] = None) -> dict[str, Any]:
    """The attempts that pass the filters, newest first, plus the choices the
    filter controls offer (taken from ALL attempts, not the filtered ones, so
    narrowing one filter never empties another's list)."""
    course_names = course_names or {}
    everything = attempt_index(conn)
    shown = [a for a in everything if flt.matches(a)]
    shown.sort(key=lambda a: (a["submitted_at"] or "", a["attempt_id"]), reverse=True)
    tries: dict[tuple[str, str], int] = defaultdict(int)
    for a in everything:
        if a["candidate_id"]:
            tries[(a["candidate_id"], a["course_id"])] += 1
    days = sorted(a["day"] for a in everything if a["day"])
    marked = [a for a in shown if a["passed"] is not None]
    pcts = [a["percentage"] for a in shown if a["percentage"] is not None]
    return {
        "attempts": [{
            "attemptId": a["attempt_id"],
            "candidateId": a["candidate_id"],
            "courseId": a["course_id"],
            "course": course_names.get(a["course_id"]) or a["course_tab"] or a["course_id"],
            "role": a["role"],
            "company": a["company"],
            "submittedAt": a["submitted_at"],
            "day": a["day"],
            "attemptNo": a["attempt_no"],
            "ofAttempts": tries.get((a["candidate_id"], a["course_id"])) if a["candidate_id"] else None,
            "score": a["score"],
            "total": a["total"],
            "percentage": a["percentage"],
            "passed": None if a["passed"] is None else bool(a["passed"]),
        } for a in shown],
        "summary": {
            "attempts": len(shown),
            "candidates": len({a["candidate_id"] for a in shown if a["candidate_id"]}),
            "passed": sum(1 for a in marked if a["passed"]),
            "failed": sum(1 for a in marked if not a["passed"]),
            "meanPercentage": round(sum(pcts) / len(pcts), 3) if pcts else None,
        },
        "options": {
            "courses": [{"id": c, "name": course_names.get(c) or c}
                        for c in sorted({a["course_id"] for a in everything if a["course_id"]})],
            "candidates": sorted({a["candidate_id"] for a in everything if a["candidate_id"]}),
            "maxAttemptNo": max(tries.values(), default=0),
            "firstDay": days[0] if days else "",
            "lastDay": days[-1] if days else "",
        },
    }

