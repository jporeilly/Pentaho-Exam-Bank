"""PowerPoint reader for extracting slides and speaker notes."""

from pathlib import Path
from typing import List, Optional
from pptx import Presentation

# SlideInfo is the generic source-section type and now lives in `source`.
# Re-exported here so `from .pptx_reader import SlideInfo` keeps working while
# this module is on its way out.
from .source import SlideInfo

__all__ = ["SlideInfo", "PPTXReader"]


class PPTXReader:
    """Reads PowerPoint files and extracts slide information."""

    def __init__(self, pptx_path: Path):
        self.pptx_path = Path(pptx_path)
        self._presentation = None
        self._slides_info: List[SlideInfo] = []

    def load(self) -> bool:
        try:
            self._presentation = Presentation(str(self.pptx_path))
            self._extract_slides_info()
            return True
        except Exception as e:
            print(f"Error loading PowerPoint: {e}")
            return False

    def _extract_slides_info(self):
        self._slides_info = []
        for idx, slide in enumerate(self._presentation.slides):
            notes = ""
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes = notes_slide.notes_text_frame.text.strip()

            title = None
            if slide.shapes.title:
                title = slide.shapes.title.text

            body_parts = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        body_parts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_text:
                            body_parts.append(" | ".join(row_text))
            body_text = "\n".join(body_parts) if body_parts else None

            self._slides_info.append(SlideInfo(
                index=idx, speaker_notes=notes, title=title, body_text=body_text,
            ))

    @property
    def slide_count(self) -> int:
        return len(self._slides_info)

    @property
    def slides(self) -> List[SlideInfo]:
        return self._slides_info

    def get_slide(self, index: int) -> Optional[SlideInfo]:
        if 0 <= index < len(self._slides_info):
            return self._slides_info[index]
        return None

    def get_speaker_notes(self, index: int) -> str:
        slide = self.get_slide(index)
        return slide.speaker_notes if slide else ""

    def has_speaker_notes(self) -> bool:
        return any(slide.speaker_notes for slide in self._slides_info)
