"""The results workbook: every tab's attempts, repeats kept apart, filters."""
from __future__ import annotations

import base64
import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from exam_bank.core import item_analysis as ia
from exam_bank.core import results_workbook as rw
from exam_bank.core.bank import ExamBankDB

ATTEMPT_HEAD = ["Attempt ID", "Candidate ID", "Course ID", "Role", "Submitted at", "Name", "Email",
                "Company", "Notes", "Score", "Total", "Percentage", "Passed", "Partner ID"]
DETAIL_HEAD = ["Attempt ID", "Course ID", "Submitted at", "Email", "Question ID", "Module", "Prompt",
               "User answer", "Correct answer", "Was correct?"]

DI, BA = "developer-di-practitioner", "analyst-ba-practitioner"


def attempt(aid, cand, course, when, score, passed, total=40):
    return [aid, cand, course, "developer", when, "Test Person", "test@example.invalid", "Acme",
            "a note", score, total, round(score / total, 2), passed, "p1"]


A1 = attempt("a1", "cand_a", DI, datetime(2026, 10, 13, 20, 3), 33, "PASS")
A2 = attempt("a2", "cand_a", DI, datetime(2026, 10, 16, 16, 34), 20, "FAIL")  # same candidate again
A3 = attempt("a3", "cand_b", DI, datetime(2026, 10, 20, 9, 0), 40, "PASS")
B1 = attempt("b1", "cand_a", BA, datetime(2026, 10, 14, 8, 0), 30, "PASS")


def workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "All results"
    for r in (ATTEMPT_HEAD, A1, A2, A3, B1):
        ws.append(r)
    for title, rows in (("DI Practitioner", [A1, A2, A3]), ("BA Practitioner", [B1]), ("SD Specialty", [])):
        t = wb.create_sheet(title)
        for r in (ATTEMPT_HEAD, *rows):
            t.append(r)
    d = wb.create_sheet("Detail")
    d.append(DETAIL_HEAD)
    for aid, ok in (("a1", "Y"), ("a2", "N")):
        d.append([aid, DI, datetime(2026, 10, 13), "test@example.invalid", "di-m1-q1", "M", "P?", "x", "y", ok])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def db(tmp_path):
    d = ExamBankDB(tmp_path / "w.db")
    yield d
    d.close()


def test_every_course_tab_is_read_and_each_attempt_once():
    parsed = rw.parse_workbook(workbook())
    assert sorted(a["attempt_id"] for a in parsed["attempts"]) == ["a1", "a2", "a3", "b1"]
    kinds = {s["name"]: s["kind"] for s in parsed["sheets"]}
    assert kinds == {"All results": "attempts", "DI Practitioner": "attempts",
                     "BA Practitioner": "attempts", "SD Specialty": "attempts", "Detail": "answers"}
    # The course tab labels the attempt; the repeated "All results" row does not override it.
    assert {a["attempt_id"]: a["course_tab"] for a in parsed["attempts"]}["b1"] == "BA Practitioner"
    assert len(parsed["answers"]) == 2


def test_personal_data_is_never_read():
    blob = repr(rw.parse_workbook(workbook()))
    assert "test@example.invalid" not in blob and "Test Person" not in blob and "a note" not in blob


def test_a_blank_course_id_falls_back_to_the_tab_name():
    wb = Workbook()
    ws = wb.active
    ws.title = "SD Specialty"
    ws.append(ATTEMPT_HEAD)
    ws.append(attempt("s1", "cand_z", "", datetime(2026, 10, 1), 10, "FAIL"))
    buf = io.BytesIO()
    wb.save(buf)
    assert rw.parse_workbook(buf.getvalue())["attempts"][0]["course_id"] == "SD Specialty"


def test_not_a_results_workbook_is_refused():
    wb = Workbook()
    wb.active.append(["Name", "Score"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(rw.WorkbookError, match="not the results workbook"):
        rw.parse_workbook(buf.getvalue())
    with pytest.raises(rw.WorkbookError, match="not an Excel"):
        rw.parse_workbook(b"Name,Score\na,1\n")


def test_repeat_attempts_stay_separate_and_are_numbered(db):
    rw.import_workbook(db.conn, workbook())
    shown = {a["attemptId"]: a for a in rw.list_attempts(db.conn, rw.Filters())["attempts"]}
    assert (shown["a1"]["attemptNo"], shown["a2"]["attemptNo"]) == (1, 2)
    assert shown["a1"]["ofAttempts"] == 2 and shown["a3"]["attemptNo"] == 1
    # Same candidate, different course: that is attempt 1 of the other course.
    assert shown["b1"]["attemptNo"] == 1 and shown["b1"]["ofAttempts"] == 1


def test_import_is_idempotent_and_adds_only_new_attempts(db):
    first = rw.import_workbook(db.conn, workbook())
    again = rw.import_workbook(db.conn, workbook())
    assert first["addedAttempts"] == 4 and again["addedAttempts"] == 0
    assert again["totalAttempts"] == 4 and again["totalResponses"] == 2
    assert first["courses"] == sorted([DI, BA])


def test_filters(db):
    rw.import_workbook(db.conn, workbook())

    def ids(**kw):
        return [a["attemptId"] for a in rw.list_attempts(db.conn, rw.Filters(**kw))["attempts"]]

    assert ids() == ["a3", "a2", "b1", "a1"]  # newest first
    assert ids(course=DI) == ["a3", "a2", "a1"]
    assert ids(candidate="cand_a") == ["a2", "b1", "a1"]
    assert ids(passed="fail") == ["a2"]
    assert ids(passed="pass", course=DI) == ["a3", "a1"]
    assert ids(attempt_no=2) == ["a2"]
    assert ids(date_from="2026-10-14", date_to="2026-10-16") == ["a2", "b1"]  # inclusive both ends
    assert ids(date_from="2026-10-17") == ["a3"]
    assert ids(date_to="2026-10-13") == ["a1"]
    with pytest.raises(ValueError):
        rw.Filters(date_from="13/10/2026")
    with pytest.raises(ValueError):
        rw.Filters(passed="maybe")


def test_the_summary_follows_the_filters_but_the_options_do_not(db):
    rw.import_workbook(db.conn, workbook())
    r = rw.list_attempts(db.conn, rw.Filters(course=BA), {DI: "DI Practitioner"})
    assert r["summary"]["attempts"] == 1 and r["summary"]["passed"] == 1
    assert [c["id"] for c in r["options"]["courses"]] == [BA, DI]
    assert r["options"]["candidates"] == ["cand_a", "cand_b"]
    assert r["options"]["maxAttemptNo"] == 2
    assert (r["options"]["firstDay"], r["options"]["lastDay"]) == ("2026-10-13", "2026-10-20")


def test_a_csv_only_attempt_still_filters_by_date_and_course(db):
    csv_text = ("Attempt ID,Course ID,Question ID,User answer,Correct answer,Was correct?,Submitted at\n"
                f"c1,{DI},q1,x,x,Y,2026-10-05T10:00:00Z\n")
    ia.import_rows(db.conn, ia.parse_detail_csv(csv_text))
    flt = rw.Filters(course=DI, date_from="2026-10-05", date_to="2026-10-05")
    assert rw.matching_attempt_ids(db.conn, flt) == {"c1"}
    assert rw.matching_attempt_ids(db.conn, rw.Filters(candidate="cand_a")) == set()
    assert rw.matching_attempt_ids(db.conn, rw.Filters()) is None


def test_analysis_follows_the_filters(db):
    rw.import_workbook(db.conn, workbook())
    assert ia.analyse(db)["exams"][0]["attempts"] == 2
    only = rw.matching_attempt_ids(db.conn, rw.Filters(passed="pass"))
    assert ia.analyse(db, attempt_ids=only)["exams"][0]["attempts"] == 1
    assert ia.analyse(db, attempt_ids=rw.matching_attempt_ids(db.conn, rw.Filters(course=BA)))["exams"] == []


def test_the_api(tmp_path):
    from exam_bank.api.app import app
    from exam_bank.api.deps import get_db

    path = tmp_path / "api.db"
    ExamBankDB(path).close()

    def override():
        d = ExamBankDB(path)
        try:
            yield d
        finally:
            d.close()

    app.dependency_overrides[get_db] = override
    try:
        c = TestClient(app)
        assert c.post("/api/results/import-xlsx", json={"xlsx": base64.b64encode(b"nope").decode()}).status_code == 422
        r = c.post("/api/results/import-xlsx", json={"xlsx": base64.b64encode(workbook()).decode()})
        assert r.status_code == 200
        body = r.json()
        assert body["attempts"] == 4 and body["answers"] == 2 and len(body["sheets"]) == 5
        got = c.get("/api/results/attempts", params={"course": DI, "attempt_no": 2}).json()
        assert [a["attemptId"] for a in got["attempts"]] == ["a2"]
        got = c.get("/api/results/attempts", params={"date_from": "2026-10-16", "date_to": "2026-10-20"}).json()
        assert [a["attemptId"] for a in got["attempts"]] == ["a3", "a2"]
        assert c.get("/api/results/attempts", params={"date_from": "soon"}).status_code == 422
        assert c.get("/api/results/analysis", params={"passed": "fail"}).json()["exams"][0]["attempts"] == 1
    finally:
        app.dependency_overrides.clear()
