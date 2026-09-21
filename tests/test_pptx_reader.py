"""Tests for core/pptx_reader.py — SlideInfo dataclass and PPTXReader."""

import pytest
from pathlib import Path

from exam_bank.core.pptx_reader import SlideInfo, PPTXReader


class TestSlideInfo:
    def test_basic_creation(self):
        slide = SlideInfo(index=0, speaker_notes="Test notes")
        assert slide.index == 0
        assert slide.speaker_notes == "Test notes"
        assert slide.title is None
        assert slide.body_text is None

    def test_full_creation(self):
        slide = SlideInfo(
            index=5,
            speaker_notes="Notes here",
            title="My Title",
            body_text="Body content",
            thumbnail_path=Path("thumb.png"),
        )
        assert slide.title == "My Title"
        assert slide.body_text == "Body content"
        assert slide.thumbnail_path == Path("thumb.png")


class TestPPTXReaderWithoutFile:
    def test_load_nonexistent(self, tmp_path):
        reader = PPTXReader(tmp_path / "nonexistent.pptx")
        assert not reader.load()
        assert reader.slide_count == 0

    def test_get_slide_out_of_range(self, tmp_path):
        reader = PPTXReader(tmp_path / "nonexistent.pptx")
        assert reader.get_slide(0) is None
        assert reader.get_slide(-1) is None

    def test_get_speaker_notes_empty(self, tmp_path):
        reader = PPTXReader(tmp_path / "nonexistent.pptx")
        assert reader.get_speaker_notes(0) == ""

    def test_has_speaker_notes_empty(self, tmp_path):
        reader = PPTXReader(tmp_path / "nonexistent.pptx")
        assert not reader.has_speaker_notes()


class TestPPTXReaderWithFile:
    """Tests that require a real PPTX file — create a minimal one with python-pptx."""

    @pytest.fixture
    def sample_pptx(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        # Slide 1: with title and notes
        layout = prs.slide_layouts[0]  # Title Slide
        slide1 = prs.slides.add_slide(layout)
        slide1.shapes.title.text = "Slide One Title"
        # Add speaker notes
        notes_slide = slide1.notes_slide
        notes_slide.notes_text_frame.text = "These are the speaker notes for slide 1."
        # Add body text
        body = slide1.placeholders[1]
        body.text = "Body text content"

        # Slide 2: with notes but no title
        layout2 = prs.slide_layouts[6]  # Blank
        slide2 = prs.slides.add_slide(layout2)
        notes2 = slide2.notes_slide
        notes2.notes_text_frame.text = "Notes for slide 2."

        # Slide 3: no notes
        slide3 = prs.slides.add_slide(layout2)

        path = tmp_path / "test.pptx"
        prs.save(str(path))
        return path

    def test_load_success(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        assert reader.load()
        assert reader.slide_count == 3

    def test_slide_titles(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        assert reader.slides[0].title == "Slide One Title"
        assert reader.slides[1].title is None

    def test_speaker_notes(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        assert "speaker notes for slide 1" in reader.slides[0].speaker_notes
        assert "Notes for slide 2" in reader.slides[1].speaker_notes
        assert reader.slides[2].speaker_notes == ""

    def test_has_speaker_notes(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        assert reader.has_speaker_notes()

    def test_get_slide(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        slide = reader.get_slide(0)
        assert slide is not None
        assert slide.index == 0

    def test_get_speaker_notes(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        notes = reader.get_speaker_notes(0)
        assert "speaker notes" in notes

    def test_body_text(self, sample_pptx):
        reader = PPTXReader(sample_pptx)
        reader.load()
        assert reader.slides[0].body_text is not None
        assert "Body text content" in reader.slides[0].body_text
