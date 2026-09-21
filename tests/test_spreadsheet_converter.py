"""Tests for core/spreadsheet_converter.py — Excel to Exam Bank CSV conversion."""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from exam_bank.core.spreadsheet_converter import (
    match_headers, convert_sheet, convert_workbook, convert_xlsx_to_csv,
    import_from_xlsx, _normalize_header, _resolve_correct_indicator,
    _detect_layout, CSV_COLUMNS,
)


# ── Helpers ───────────────────────────────────────────


def _make_workbook(tmp_path, sheets: dict) -> Path:
    """Create a .xlsx file from a dict of {sheet_name: [rows]}.

    The first row in each sheet is the header.
    """
    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)

    path = tmp_path / "test.xlsx"
    wb.save(path)
    wb.close()
    return path


# ── Header normalization ──────────────────────────────


class TestNormalizeHeader:
    def test_simple(self):
        assert _normalize_header("Stem") == "stem"

    def test_spaces_to_underscores(self):
        assert _normalize_header("Bloom Level") == "bloom_level"

    def test_special_chars(self):
        assert _normalize_header("Bloom's Level") == "blooms_level"

    def test_strip(self):
        assert _normalize_header("  Topic  ") == "topic"


# ── Header matching ───────────────────────────────────


class TestMatchHeaders:
    def test_direct_columns(self):
        headers = ["Question", "Answer", "Distractor 1", "Distractor 2", "Topic"]
        mapping = match_headers(headers)
        assert mapping["stem"] == 0
        assert mapping["key"] == 1
        assert mapping["distractor_1"] == 2
        assert mapping["distractor_2"] == 3
        assert mapping["topic"] == 4

    def test_option_columns(self):
        headers = ["Question", "Option A", "Option B", "Option C", "Option D", "Correct"]
        mapping = match_headers(headers)
        assert mapping["stem"] == 0
        assert mapping["option_a"] == 1
        assert mapping["option_b"] == 2
        assert mapping["correct_indicator"] == 5

    def test_choice_columns(self):
        headers = ["Stem", "Choice 1", "Choice 2", "Choice 3", "Choice 4", "Answer Key"]
        mapping = match_headers(headers)
        assert mapping["stem"] == 0
        assert mapping["option_a"] == 1  # choice_1 maps to option_a
        assert mapping["correct_indicator"] == 5

    def test_fuzzy_matching(self):
        headers = ["Question Text", "Right Answer", "Explanation"]
        mapping = match_headers(headers)
        assert "stem" in mapping
        assert "explanation" in mapping


# ── Layout detection ──────────────────────────────────


class TestDetectLayout:
    def test_direct_layout(self):
        mapping = {"stem": 0, "key": 1, "distractor_1": 2}
        assert _detect_layout(mapping) == "direct"

    def test_options_layout(self):
        mapping = {"stem": 0, "option_a": 1, "option_b": 2, "correct_indicator": 3}
        assert _detect_layout(mapping) == "options"


# ── Correct indicator resolution ──────────────────────


class TestResolveCorrectIndicator:
    def test_letter_indicator(self):
        options = {"option_a": "TCP", "option_b": "UDP", "option_c": "HTTP", "option_d": "FTP"}
        key, distractors = _resolve_correct_indicator("B", options)
        assert key == "UDP"
        assert "TCP" in distractors
        assert "UDP" not in distractors

    def test_number_indicator(self):
        options = {"option_a": "TCP", "option_b": "UDP", "option_c": "HTTP"}
        key, distractors = _resolve_correct_indicator("3", options)
        assert key == "HTTP"

    def test_lowercase_letter(self):
        options = {"option_a": "TCP", "option_b": "UDP"}
        key, _ = _resolve_correct_indicator("a", options)
        assert key == "TCP"

    def test_text_match(self):
        options = {"option_a": "TCP", "option_b": "UDP"}
        key, distractors = _resolve_correct_indicator("UDP", options)
        assert key == "UDP"
        assert distractors == ["TCP"]

    def test_option_prefix(self):
        options = {"option_a": "TCP", "option_b": "UDP", "option_c": "HTTP"}
        key, _ = _resolve_correct_indicator("Option C", options)
        assert key == "HTTP"


# ── Sheet conversion — Direct pattern ─────────────────


class TestConvertSheetDirect:
    def test_basic_conversion(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2", "Distractor 3"])
        ws.append(["What is HTTP?", "HyperText Transfer Protocol", "File Transfer", "Simple Mail", "Telnet"])
        ws.append(["What is DNS?", "Domain Name System", "Data Network", "Digital Name", ""])

        rows = convert_sheet(ws, sheet_name="Networking", default_topic="Networking")
        assert len(rows) == 2
        assert rows[0]["Stem"] == "What is HTTP?"
        assert rows[0]["Key"] == "HyperText Transfer Protocol"
        assert rows[0]["Distractor 1"] == "File Transfer"
        assert rows[0]["Topic"] == "Networking"  # from default_topic (sheet name)

    def test_with_topic_column(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2", "Topic"])
        ws.append(["Q?", "A", "B", "C", "Security"])

        rows = convert_sheet(ws, default_topic="Fallback")
        assert rows[0]["Topic"] == "Security"  # column takes precedence

    def test_blank_rows_skipped(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2"])
        ws.append(["Q1?", "A1", "B1", "C1"])
        ws.append([None, None, None, None])  # blank row
        ws.append(["Q2?", "A2", "B2", "C2"])

        rows = convert_sheet(ws)
        assert len(rows) == 2

    def test_missing_stem_skipped(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1"])
        ws.append(["", "A1", "B1"])  # no stem
        ws.append(["Q2?", "A2", "B2"])

        rows = convert_sheet(ws)
        assert len(rows) == 1
        assert rows[0]["Stem"] == "Q2?"

    def test_pipe_separated_multi_key(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Key", "Distractor 1", "Distractor 2"])
        ws.append(["Which two? (Select 2)", "Ans A|Ans B", "Wrong 1", "Wrong 2"])

        rows = convert_sheet(ws)
        assert rows[0]["Question Type"] == "multi"
        assert rows[0]["Keys"] == "Ans A|Ans B"
        assert rows[0]["Key"] == "Ans A"

    def test_difficulty_and_bloom(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Stem", "Key", "Distractor 1", "Distractor 2", "Difficulty", "Bloom Level"])
        ws.append(["Q?", "A", "B", "C", "Hard", "Analyze"])

        rows = convert_sheet(ws)
        assert rows[0]["Difficulty"] == "Hard"
        assert rows[0]["Bloom Level"] == "Analyze"

    def test_defaults_for_missing_difficulty_bloom(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2"])
        ws.append(["Q?", "A", "B", "C"])

        rows = convert_sheet(ws)
        assert rows[0]["Difficulty"] == "Medium"
        assert rows[0]["Bloom Level"] == "Apply"


# ── Sheet conversion — Option pattern ─────────────────


class TestConvertSheetOptions:
    def test_option_abcd_with_letter(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Option A", "Option B", "Option C", "Option D", "Correct"])
        ws.append(["What is 2+2?", "3", "4", "5", "6", "B"])

        rows = convert_sheet(ws)
        assert len(rows) == 1
        assert rows[0]["Key"] == "4"
        assert "3" in [rows[0]["Distractor 1"], rows[0]["Distractor 2"], rows[0]["Distractor 3"]]

    def test_option_with_number_indicator(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Choice 1", "Choice 2", "Choice 3", "Correct"])
        ws.append(["Capital of France?", "London", "Paris", "Berlin", "2"])

        rows = convert_sheet(ws)
        assert rows[0]["Key"] == "Paris"

    def test_five_options(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Option A", "Option B", "Option C", "Option D", "Option E", "Correct"])
        ws.append(["Q?", "A1", "A2", "A3", "A4", "A5", "E"])

        rows = convert_sheet(ws)
        assert rows[0]["Key"] == "A5"
        assert len([rows[0][f"Distractor {i}"] for i in range(1, 5) if rows[0][f"Distractor {i}"]]) == 4


# ── Multi-sheet workbook ──────────────────────────────


class TestConvertWorkbook:
    def test_combined(self, tmp_path):
        path = _make_workbook(tmp_path, {
            "Networking": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q1?", "A1", "B1", "C1"],
            ],
            "Security": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q2?", "A2", "B2", "C2"],
                ["Q3?", "A3", "B3", "C3"],
            ],
        })

        rows = convert_workbook(path, combine=True)
        assert len(rows) == 3
        topics = [r["Topic"] for r in rows]
        assert "Networking" in topics
        assert "Security" in topics

    def test_per_sheet(self, tmp_path):
        path = _make_workbook(tmp_path, {
            "Tab1": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q1?", "A1", "B1", "C1"],
            ],
            "Tab2": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q2?", "A2", "B2", "C2"],
            ],
        })

        result = convert_workbook(path, combine=False)
        assert "Tab1" in result
        assert "Tab2" in result
        assert len(result["Tab1"]) == 1
        assert len(result["Tab2"]) == 1

    def test_skips_instructions_sheet(self, tmp_path):
        path = _make_workbook(tmp_path, {
            "Instructions": [
                ["This sheet contains instructions"],
                ["Do not delete this sheet"],
            ],
            "Questions": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q?", "A", "B", "C"],
            ],
        })

        rows = convert_workbook(path, combine=True)
        assert len(rows) == 1

    def test_sheet_name_as_topic_disabled(self, tmp_path):
        path = _make_workbook(tmp_path, {
            "MySheet": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q?", "A", "B", "C"],
            ],
        })

        rows = convert_workbook(path, combine=True, use_sheet_as_topic=False)
        assert rows[0]["Topic"] == ""


# ── CSV output ────────────────────────────────────────


class TestCSVOutput:
    def test_convert_to_csv(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "Sheet1": [
                ["Question", "Answer", "Distractor 1", "Distractor 2", "Distractor 3"],
                ["What is HTTP?", "Protocol", "Format", "Language", "Tool"],
            ],
        })

        paths = convert_xlsx_to_csv(xlsx, output_path=tmp_path / "output.csv")
        assert len(paths) == 1
        assert paths[0].exists()

        # Read back and verify
        with open(paths[0], encoding="utf-8") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        assert len(rows) == 1
        assert rows[0]["Stem"] == "What is HTTP?"

    def test_backup_copy(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "Sheet1": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q?", "A", "B", "C"],
            ],
        })
        backup_dir = tmp_path / "backups"

        convert_xlsx_to_csv(xlsx, output_path=tmp_path / "out.csv", backup_dir=backup_dir)
        assert (backup_dir / "out.csv").exists()

    def test_per_sheet_csv(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "Tab1": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q1?", "A1", "B1", "C1"],
            ],
            "Tab2": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q2?", "A2", "B2", "C2"],
            ],
        })

        paths = convert_xlsx_to_csv(xlsx, output_path=tmp_path / "out.csv", combine=False)
        assert len(paths) == 2


# ── Round-trip through importer ───────────────────────


class TestImportFromXlsx:
    def test_roundtrip(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "Security": [
                ["Question", "Answer", "Distractor 1", "Distractor 2", "Distractor 3",
                 "Explanation", "Difficulty"],
                ["What is HTTPS?", "Secure HTTP", "FTP", "SMTP", "DNS",
                 "HTTPS uses TLS encryption", "Hard"],
            ],
        })

        questions = import_from_xlsx(xlsx)
        assert len(questions) == 1
        q = questions[0]
        assert q.stem == "What is HTTPS?"
        assert q.key == "Secure HTTP"
        assert len(q.distractors) == 3
        assert q.explanation == "HTTPS uses TLS encryption"
        assert q.difficulty == "Hard"
        assert q.topic == "Security"  # from sheet name

    def test_backup_created(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "Sheet1": [
                ["Question", "Answer", "Distractor 1", "Distractor 2"],
                ["Q?", "A", "B", "C"],
            ],
        })
        backup_dir = tmp_path / "backup"

        import_from_xlsx(xlsx, backup_dir=backup_dir)
        backup_files = list(backup_dir.glob("*.csv"))
        assert len(backup_files) == 1

    def test_option_pattern_roundtrip(self, tmp_path):
        xlsx = _make_workbook(tmp_path, {
            "General": [
                ["Question", "Option A", "Option B", "Option C", "Option D", "Correct", "Topic"],
                ["What color is the sky?", "Red", "Blue", "Green", "Yellow", "B", "Science"],
            ],
        })

        questions = import_from_xlsx(xlsx)
        assert len(questions) == 1
        q = questions[0]
        assert q.key == "Blue"
        assert "Red" in q.distractors
        assert q.topic == "Science"  # from column, not sheet name

    def test_empty_workbook(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Empty"
        path = tmp_path / "empty.xlsx"
        wb.save(path)
        wb.close()

        questions = import_from_xlsx(path)
        assert questions == []


# ── Merged cells ──────────────────────────────────────


class TestMergedCells:
    def test_merged_scenario(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2", "Scenario"])
        ws.append(["Q1?", "A1", "B1", "C1", "Shared scenario"])
        ws.append(["Q2?", "A2", "B2", "C2", None])
        # Merge the scenario cells
        ws.merge_cells("E2:E3")

        rows = convert_sheet(ws)
        assert len(rows) == 2
        assert rows[0]["Scenario"] == "Shared scenario"
        assert rows[1]["Scenario"] == "Shared scenario"


# ── Multi-select detection from stem ──────────────────


class TestMultiSelectDetection:
    def test_which_two_detected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2"])
        ws.append(["Which two protocols are secure? (Select 2)", "TLS|SSH", "FTP", "Telnet"])

        rows = convert_sheet(ws)
        assert rows[0]["Question Type"] == "multi"

    def test_choose_n_detected(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Question", "Answer", "Distractor 1", "Distractor 2"])
        ws.append(["Select the best options. (Choose 3)", "A", "B", "C"])

        rows = convert_sheet(ws)
        assert rows[0]["Question Type"] == "multi"
